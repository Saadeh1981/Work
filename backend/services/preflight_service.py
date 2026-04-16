from __future__ import annotations

import os
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from concurrent.futures import ThreadPoolExecutor, as_completed

import fitz  # PyMuPDF
from PIL import Image

try:
    import docx
except Exception:
    docx = None

try:
    import openpyxl
except Exception:
    openpyxl = None


@dataclass
class DocProfile:
    path: str
    name: str
    ext: str
    size_bytes: int
    kind: str

    page_count: Optional[int] = None
    sheet_count: Optional[int] = None
    encrypted: Optional[bool] = None
    has_text_layer: Optional[bool] = None
    avg_text_chars_sampled: Optional[int] = None
    avg_images_sampled: Optional[float] = None
    scan_likelihood: Optional[float] = None

    warnings: List[str] = None
    plan: str = "unknown"

    def to_dict(self) -> Dict[str, Any]:
        d = asdict(self)
        d["warnings"] = self.warnings or []
        return d


def _safe_int(v: Any) -> Optional[int]:
    try:
        return int(v)
    except Exception:
        return None


def scan_folder(root: str, max_files: int = 500) -> List[str]:
    p = Path(root)
    if p.is_file():
        if p.name.startswith("~$"):
            return []
        return [str(p)]

    allowed = {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".docx", ".xlsx"}
    files: List[str] = []

    for f in p.rglob("*"):
        if not f.is_file():
            continue
        if f.name.startswith("~$"):
            continue
        if f.suffix.lower() in allowed:
            files.append(str(f))
        if len(files) >= max_files:
            break

    return files


def preflight(paths: List[str], sample_pages: int = 5) -> List[DocProfile]:
    results: List[DocProfile] = []
    with ThreadPoolExecutor(max_workers=min(16, max(1, len(paths)))) as ex:
        futs = {ex.submit(preflight_one, p, sample_pages): p for p in paths}
        for fut in as_completed(futs):
            results.append(fut.result())
    results.sort(key=lambda x: (x.kind, -(x.page_count or 0), -x.size_bytes))
    return results


def preflight_one(path: str, sample_pages: int = 5) -> DocProfile:
    p = Path(path)
    ext = p.suffix.lower()
    size_bytes = p.stat().st_size if p.exists() else 0

    prof = DocProfile(
        path=str(p),
        name=p.name,
        ext=ext,
        size_bytes=size_bytes,
        kind=_classify_kind(ext),
        warnings=[],
    )

    if prof.kind == "pdf":
        _preflight_pdf(prof, sample_pages)
    elif prof.kind == "image":
        _preflight_image(prof)
    elif prof.kind == "docx":
        _preflight_docx(prof)
    elif prof.kind == "xlsx":
        _preflight_xlsx(prof)
    else:
        prof.warnings.append("Unsupported file type")
        prof.plan = "skip"

    _choose_plan(prof)
    return prof


def _classify_kind(ext: str) -> str:
    if ext == ".pdf":
        return "pdf"
    if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        return "image"
    if ext == ".docx":
        return "docx"
    if ext == ".xlsx":
        return "xlsx"
    return "other"


def _preflight_pdf(prof: DocProfile, sample_pages: int) -> None:
    try:
        doc = fitz.open(prof.path)
    except Exception as e:
        prof.warnings.append(f"PDF open failed: {type(e).__name__}")
        prof.plan = "skip"
        return

    prof.page_count = doc.page_count
    prof.encrypted = bool(getattr(doc, "is_encrypted", False))

    if prof.encrypted:
        prof.warnings.append("PDF is encrypted")
        doc.close()
        return

    pages_to_check = list(range(min(sample_pages, doc.page_count)))
    text_chars = []
    image_counts = []

    for i in pages_to_check:
        try:
            page = doc.load_page(i)
            txt = page.get_text("text") or ""
            text_chars.append(len(txt.strip()))
            images = page.get_images(full=True)
            image_counts.append(len(images))
        except Exception:
            prof.warnings.append(f"PDF page {i + 1} read failed")
            continue

    doc.close()

    if text_chars:
        prof.avg_text_chars_sampled = int(sum(text_chars) / len(text_chars))
    if image_counts:
        prof.avg_images_sampled = float(sum(image_counts) / len(image_counts))

    prof.has_text_layer = bool((prof.avg_text_chars_sampled or 0) > 200)

    prof.scan_likelihood = _estimate_scan_likelihood(
        avg_text_chars=prof.avg_text_chars_sampled or 0,
        avg_images=prof.avg_images_sampled or 0.0,
        page_count=prof.page_count or 0,
        size_bytes=prof.size_bytes,
    )

    if (prof.page_count or 0) > 400:
        prof.warnings.append("Huge PDF, enable slicing")
    if (prof.size_bytes / (1024 * 1024)) > 200:
        prof.warnings.append("Large PDF size, expect slow rasterization")


def _estimate_scan_likelihood(
    avg_text_chars: int,
    avg_images: float,
    page_count: int,
    size_bytes: int,
) -> float:
    score = 0.0
    if avg_text_chars < 50:
        score += 0.6
    elif avg_text_chars < 200:
        score += 0.3
    if avg_images >= 2:
        score += 0.3
    if page_count >= 200:
        score += 0.1
    if size_bytes >= 50 * 1024 * 1024:
        score += 0.1
    if score > 1.0:
        score = 1.0
    return score


def _preflight_image(prof: DocProfile) -> None:
    try:
        with Image.open(prof.path) as im:
            w, h = im.size
        prof.page_count = 1
        prof.has_text_layer = False
        prof.scan_likelihood = 1.0
        if w < 900 or h < 900:
            prof.warnings.append("Low resolution image, OCR quality risk")
    except Exception as e:
        prof.warnings.append(f"Image open failed: {type(e).__name__}")
        prof.plan = "skip"


def _preflight_docx(prof: DocProfile) -> None:
    if docx is None:
        prof.warnings.append("python-docx not installed")
        return
    try:
        d = docx.Document(prof.path)
        para_count = len(d.paragraphs)
        prof.page_count = None
        prof.has_text_layer = True
        prof.scan_likelihood = 0.0
        if para_count == 0:
            prof.warnings.append("Empty DOCX or non text content")
    except Exception as e:
        prof.warnings.append(f"DOCX open failed: {type(e).__name__}")
        prof.plan = "skip"


def _preflight_xlsx(prof: DocProfile) -> None:
    if openpyxl is None:
        prof.warnings.append("openpyxl not installed")
        return
    try:
        wb = openpyxl.load_workbook(prof.path, read_only=True, data_only=True)
        sheet_count = len(wb.sheetnames)
        prof.page_count = None
        prof.sheet_count = sheet_count
        prof.has_text_layer = True
        prof.scan_likelihood = 0.0
        if sheet_count == 0:
            prof.warnings.append("Empty XLSX")
        wb.close()
    except Exception as e:
        prof.warnings.append(f"XLSX open failed: {type(e).__name__}")
        prof.plan = "skip"


def _choose_plan(prof: DocProfile) -> None:
    if prof.plan == "skip":
        return

    if prof.kind == "pdf":
        if prof.encrypted:
            prof.plan = "blocked_encrypted"
            return
        scan_like = prof.scan_likelihood or 0.0
        pages = prof.page_count or 0
        if scan_like >= 0.7:
            prof.plan = "ocr_pdf_sliced" if pages >= 80 else "ocr_pdf"
        else:
            prof.plan = "text_pdf_then_ocr_low_text_pages" if pages >= 20 else "text_pdf"
        return

    if prof.kind == "image":
        prof.plan = "ocr_image"
        return

    if prof.kind in {"docx", "xlsx"}:
        prof.plan = "native_parse"
        return

    prof.plan = "skip"


if __name__ == "__main__":
    root = os.environ.get("SCAN_ROOT", ".")
    files = scan_folder(root)
    profiles = preflight(files, sample_pages=5)
    for p in profiles[:20]:
        print(p.to_dict())
